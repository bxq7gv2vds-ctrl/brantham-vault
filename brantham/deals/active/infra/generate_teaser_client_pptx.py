from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("INFRA-teaser-client.pptx")

NAVY = RGBColor(36, 73, 122)
GREEN = RGBColor(96, 176, 88)
LIGHT_GREEN = RGBColor(229, 242, 225)
LIGHT_BLUE = RGBColor(226, 235, 247)
INK = RGBColor(45, 45, 45)
MID = RGBColor(95, 95, 95)
LINE = RGBColor(115, 115, 115)
WHITE = RGBColor(255, 255, 255)
OFF = RGBColor(248, 248, 246)
RED = RGBColor(182, 42, 58)

FONT = "Aptos"
FONT_TITLE = "Aptos Display"


def style(run, size=10, bold=False, italic=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def box(slide, x, y, w, h, fill=WHITE, line_color=None, radius=False):
    shp = slide.shapes.add_shape(5 if radius else 1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_color or fill
    shp.line.width = Pt(0.9 if line_color else 0)
    return shp


def text(slide, x, y, w, h, value, size=10, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT, font=FONT):
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = value
    style(r, size=size, bold=bold, italic=italic, color=color, font=font)
    return shp


def rich_line(slide, x, y, w, h, lead, body, size=8.3):
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = lead
    style(r, size=size, bold=True)
    r = p.add_run()
    r.text = " " + body
    style(r, size=size)
    return shp


def header_bar(slide, x, y, w, h, label):
    box(slide, x, y, w, h, NAVY)
    text(slide, x, y + 0.06, w, h - 0.08, label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font=FONT_TITLE)


def title(slide, value):
    text(slide, 0.28, 0.16, 6.7, 0.42, value, size=27, bold=True, color=NAVY, font=FONT_TITLE)


def footer(slide, page):
    text(slide, 10.75, 7.06, 1.4, 0.15, "Private & Confidential", size=7.4, italic=True, color=NAVY, align=PP_ALIGN.RIGHT)
    text(slide, 12.28, 7.06, 0.25, 0.15, str(page), size=8, color=MID, align=PP_ALIGN.RIGHT)


def bullet_list(slide, x, y, w, items, size=8.8, gap=0.38, color=INK):
    cy = y
    for item in items:
        text(slide, x, cy, 0.13, 0.18, "•", size=size + 1.5, color=GREEN, align=PP_ALIGN.CENTER)
        text(slide, x + 0.20, cy, w - 0.20, gap, item, size=size, color=color)
        cy += gap
    return cy


def small_table(slide, x, y, widths, row_h, headers, rows, header_fill=GREEN, body_size=7.4, header_size=7.5):
    total = sum(widths)
    box(slide, x, y, total, row_h, header_fill)
    cx = x
    for i, h in enumerate(headers):
        align = PP_ALIGN.RIGHT if i > 0 else PP_ALIGN.LEFT
        text(slide, cx + 0.04, y + 0.07, widths[i] - 0.08, row_h - 0.08, h, size=header_size, bold=True, color=WHITE, align=align)
        cx += widths[i]
    cy = y + row_h
    for row in rows:
        cx = x
        for i, val in enumerate(row):
            align = PP_ALIGN.RIGHT if i > 0 else PP_ALIGN.LEFT
            text(slide, cx + 0.04, cy + 0.055, widths[i] - 0.08, row_h - 0.06, str(val), size=body_size, bold=(i == 0), color=INK, align=align)
            cx += widths[i]
        # row separator
        sep = slide.shapes.add_shape(1, Inches(x), Inches(cy + row_h - 0.01), Inches(total), Pt(0.6))
        sep.fill.solid()
        sep.fill.fore_color.rgb = LINE
        sep.line.color.rgb = LINE
        cy += row_h
    return cy


def overview_rows(slide, x, y, w, rows, row_h=0.53):
    cy = y
    for icon, label, val, accent in rows:
        text(slide, x, cy + 0.08, 0.62, 0.30, icon, size=17, bold=True, color=accent, align=PP_ALIGN.CENTER)
        text(slide, x + 0.84, cy + 0.12, 2.45, 0.24, label, size=8.9, bold=True)
        text(slide, x + 3.40, cy + 0.10, w - 3.45, 0.30, val, size=8.25)
        sep = slide.shapes.add_shape(1, Inches(x), Inches(cy + row_h - 0.02), Inches(w), Pt(0.7))
        sep.fill.solid()
        sep.fill.fore_color.rgb = LINE
        sep.line.color.rgb = LINE
        cy += row_h
    return cy


def note_box(slide, x, y, w, h, value, fill=LIGHT_BLUE, border=GREEN, size=8.7):
    box(slide, x, y, w, h, fill, border, radius=True)
    text(slide, x + 0.12, y + 0.10, w - 0.24, h - 0.18, value, size=size)


def dotted_divider(slide, x, y, h):
    cy = y
    while cy < y + h:
        box(slide, x, cy, 0.025, 0.025, LINE)
        cy += 0.085


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


# Slide 1
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
title(s, "Company and Investment Overview")
box(s, 9.0, 0, 4.35, 0.58, NAVY)
footer(s, 1)

left_x, left_w = 0.22, 4.45
right_x, right_w = 4.96, 8.15
dotted_divider(s, 4.80, 0.92, 6.10)

header_bar(s, left_x, 0.92, left_w, 0.36, "Company Overview")
box(s, left_x + 0.04, 1.34, left_w - 0.08, 1.75, WHITE, GREEN, radius=True)
bullet_list(
    s,
    left_x + 0.28,
    1.56,
    left_w - 0.58,
    [
        "Agence-conseil française de communication et production visuelle, créée à la fin des années 1970.",
        "Positionnement intégré : marque, marque employeur, slow content, édition / print éco-responsable, digital frugal et Green UX.",
        "Ancrage historique Grand Est, avec savoir-faire photo / vidéo / print / digital et portefeuille B2B à documenter.",
    ],
    size=8.15,
    gap=0.44,
)

header_bar(s, left_x, 3.30, left_w, 0.36, "Financial Snapshot")
small_table(
    s,
    left_x,
    3.77,
    [1.86, 0.70, 0.70, 0.70, 0.78],
    0.29,
    ["K€", "FY22", "FY23", "FY24", "9m 25"],
    [
        ["Revenue", "2 993", "2 427", "2 312", "1 357"],
        ["EBE", "227", "36", "79", "(284)"],
        ["EBE %", "7.6%", "1.5%", "3.4%", "(20.9%)"],
        ["Net income", "97", "(77)", "(1)", "(345)"],
        ["Payroll + charges", "1 137", "1 064", "972", "820"],
    ],
    body_size=7.2,
)
text(s, left_x + 0.02, 5.64, left_w, 0.18, "Note: 9m 25 corresponds to 12/2025 management accounts; revenue annualised at c. €1.8m.", size=6.15, italic=True, color=MID)

header_bar(s, left_x, 5.94, left_w, 0.36, "Ownership / Procedure")
small_table(
    s,
    left_x,
    6.42,
    [2.65, 1.98],
    0.28,
    ["Item", "Status"],
    [
        ["Shareholder", "INTRA SAS, 100%"],
        ["Procedure", "Sauvegarde"],
        ["Transfer status", "To be clarified"],
    ],
    body_size=6.95,
    header_size=7.1,
)

header_bar(s, right_x, 0.92, right_w, 0.36, "Proposed Transaction")
note_box(
    s,
    right_x,
    1.38,
    right_w,
    0.46,
    "Opportunity to acquire and relaunch a regional communication platform, subject to confirmation of the procedural route allowing an asset transfer and the exact scope of transferable contracts, employees, IP and lease.",
    size=8.55,
)

header_bar(s, right_x, 2.02, right_w, 0.36, "Business / Asset Overview")
overview_rows(
    s,
    right_x,
    2.50,
    right_w,
    [
        ("★", "Core activity", "Brand strategy, visual production, content, print, employer branding and digital responsibility", GREEN),
        ("€", "Revenue base", "c. €2.3m last full year; c. €1.8m recent annualised revenue; historical peak c. €3.0m", GREEN),
        ("%", "Gross margin", "c. 75-77% gross margin preserved; issue lies in volume vs fixed cost base", GREEN),
        ("👥", "Manpower", "18-19 current employees; first reduction already consumed vs 21 in 2023 / 25 in legacy deck", GREEN),
        ("⚙", "Break-even", "Estimated EBE break-even around €2.4-2.6m revenue", GREEN),
        ("⌁", "Assets", "Brand, client file, photo/video library, domains, methodology and production know-how to be secured", GREEN),
        ("!", "Liabilities", "Historical PGE debt c. €850k not expected to transfer in an asset sale; L.642-12 direct charge c. €9.8k", RED),
    ],
    row_h=0.56,
)
text(s, right_x, 6.55, right_w, 0.32, "* Pre-NDA view. The exact scope remains subject to data room confirmation, procedural status and transferability of client contracts.", size=6.45, italic=True, color=MID)


# Slide 2
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
title(s, "Turnaround Diagnosis and Repreneur Thesis")
box(s, 9.0, 0, 4.35, 0.58, NAVY)
footer(s, 2)

left_x, left_w = 0.22, 4.45
right_x, right_w = 4.96, 8.15
dotted_divider(s, 4.80, 0.92, 6.10)

header_bar(s, left_x, 0.92, left_w, 0.36, "DCP Diagnosis")
box(s, left_x + 0.04, 1.36, left_w - 0.08, 2.05, WHITE, GREEN, radius=True)
bullet_list(
    s,
    left_x + 0.28,
    1.58,
    left_w - 0.60,
    [
        "LBO stress: 2016 acquisition via INTRA created a financial burden and cash upstream pressure.",
        "Commercial gap: prolonged absence of the commercial lead weakened pipeline and new business momentum.",
        "Market pressure: communication budgets were cut first in a weaker macro environment.",
        "AI shock: clients increasingly believe they can produce simple content internally.",
        "Cost base mismatch: payroll remained sized for a higher revenue base, driving negative EBE.",
    ],
    size=7.65,
    gap=0.35,
)

header_bar(s, left_x, 3.65, left_w, 0.36, "Actions Already Undertaken")
small_table(
    s,
    left_x,
    4.12,
    [2.05, 2.60],
    0.30,
    ["Action", "Read-through"],
    [
        ["Cost saving plan", "2023-25 restructuring already started"],
        ["Headcount reduction", "21 employees in 2023 -> 18-19 current"],
        ["Subcontracting", "External production partly variabilised"],
        ["Management bonuses", "Direction premiums removed"],
        ["Commercial function", "Reinforcement attempted, to be challenged"],
    ],
    body_size=6.85,
)

header_bar(s, right_x, 0.92, right_w, 0.36, "Repreneur Thesis")
note_box(
    s,
    right_x,
    1.38,
    right_w,
    0.50,
    "This is not a clean stand-alone agency acquisition. It is a relaunch of a regional platform with preserved gross margin, historical revenue capacity, a reduced employee base and identifiable break-even levers.",
    size=8.45,
)

header_bar(s, right_x, 2.08, right_w, 0.36, "Value Creation Levers")
overview_rows(
    s,
    right_x,
    2.55,
    right_w,
    [
        ("1", "Return to break-even", "Rebuild toward €2.4-2.6m revenue; recent annualised base c. €1.8m", GREEN),
        ("2", "Commercial relaunch", "Recover part of the historical c. €3.0m revenue capacity through B2B accounts", GREEN),
        ("3", "Cost recalibration", "Define a viable transferred perimeter around 18-19 employees or less, subject to project", GREEN),
        ("4", "AI / content repositioning", "Move away from commoditised production toward video, employer branding, workflows and advisory", GREEN),
        ("5", "Balance sheet reset", "Historical PGE c. €850k not expected to transfer; limited direct L.642-12 charge c. €9.8k", GREEN),
    ],
    row_h=0.53,
)

header_bar(s, right_x, 5.57, right_w, 0.36, "Key Confirmations Before Offer")
small_table(
    s,
    right_x,
    6.03,
    [1.95, 2.05, 2.05, 2.05],
    0.30,
    ["Procedure", "Contracts", "Lease / assets", "Social / IP"],
    [
        ["Sauvegarde route", "Transferable clients", "Strasbourg lease", "Updated headcount"],
        ["RJ / asset sale", "Concentration", "Sold assets list", "Paid leave / CSE"],
        ["Deadline", "Recurring revenue", "Leasing claims", "Brands / domains / file"],
    ],
    body_size=6.35,
    header_size=6.8,
)
text(s, right_x, 7.02, right_w, 0.18, "Disclaimer: preliminary and confidential teaser; not an offer, not an information memorandum, not investment advice.", size=6.1, italic=True, color=MID)

prs.save(OUT)
print(OUT)
